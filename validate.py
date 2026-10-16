"""Focused checks on the numbers behind figure.png and figure_median.png.

Run after get_data.py and plot.py. Writes small CSVs to data/processed/ and prints a
summary; validation.md explains the results.

  1. coverage    families, implicates, weight totals and definitions in every wave
  2. fed tables  reproduce the Fed's public-data Table 4 (2025 dollars) for all families,
                 the four education groups and the two family-with-children groups
  3. conventions pooled estimates (used) versus averaging implicate-specific estimates
  4. uncertainty 95% intervals for the latest wave, following the codebook's MEANIT
                 procedure: 999 replicate weights on implicate 1 for sampling variance,
                 plus (6/5) x the variance across the five implicates

Usage:  python validate.py
"""

import zipfile

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

from plot import GROUPS, PROCESSED, RAW, ROOT, load_year, survey_years, weighted_mean, weighted_median

TABLE4_ROWS = {
    "All families": lambda d: np.ones(len(d), bool),
    "No high school diploma": lambda d: d["EDCL"] == 1,
    "High school diploma": lambda d: d["EDCL"] == 2,
    "Some college": lambda d: d["EDCL"] == 3,
    "College degree": lambda d: d["EDCL"] == 4,
    "Single with child(ren)": lambda d: (d["MARRIED"] == 2) & (d["KIDS"] > 0),
    "Couple with child(ren)": lambda d: (d["MARRIED"] == 1) & (d["KIDS"] > 0),
}


def headline(group):
    return lambda d: (d["KIDS"] > 0) & d["EDCL"].isin(GROUPS[group])


def coverage(data):
    rows = []
    for year, d in data.items():
        per_fam = d.groupby("family_id")
        rows.append(dict(
            survey_year=year, records=len(d), families=d["family_id"].nunique(),
            weighted_families_millions=round(d["WGT"].sum() / 1e6, 3),
            weight_varies_across_implicates=round(float((per_fam["WGT"].nunique() > 1).mean()), 4),
            edcl_varies_across_implicates=round(float((per_fam["EDCL"].nunique() > 1).mean()), 4),
            kids_varies_across_implicates=round(float((per_fam["KIDS"].nunique() > 1).mean()), 4),
            kids_matches_famstruct=bool(((d["KIDS"] > 0) == d["FAMSTRUCT"].isin([1, 4])).all()),
            missing_records=d.attrs["n_missing_records"],
            **{f"families_with_children_{g.split()[0].lower()}": d.loc[headline(g)(d), "family_id"].nunique()
               for g in GROUPS},
        ))
    return pd.DataFrame(rows)


def fed_table4(data):
    path = RAW / f"scf{max(data)}_tables_public_real_historical.xlsx"
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["Table 4"]
    rows = list(ws.iter_rows(values_only=True))
    cols, year = {}, None
    for j, (y, stat) in enumerate(zip(rows[2], rows[3])):
        year = y if isinstance(y, int) else year
        if stat in ("Median", "Mean"):
            cols[(year, stat.lower())] = j
    out = []
    for i, r in enumerate(rows, start=1):
        label = r[0].strip() if isinstance(r[0], str) else None
        if label not in TABLE4_ROWS:
            continue
        for (y, stat), j in cols.items():
            d = data[y]
            g = d[np.asarray(TABLE4_ROWS[label](d))]
            fn = weighted_mean if stat == "mean" else weighted_median
            ours = fn(g["NETWORTH"], g["WGT"]) / 1e3
            out.append(dict(row=label, survey_year=y, statistic=stat,
                            fed_cell=f"'Table 4'!{get_column_letter(j + 1)}{i}",
                            fed_thousands=r[j], ours_thousands=round(ours, 3),
                            abs_diff_thousands=round(abs(ours - r[j]), 3),
                            rel_diff=abs(ours - r[j]) / abs(r[j])))
    return pd.DataFrame(out)


def implicate_average(g, fn):
    return float(np.mean([fn(t["NETWORTH"], t["WGT"]) for _, t in g.groupby("implicate")]))


def interpolated_median(x, w):
    x, w = np.asarray(x, float), np.asarray(w, float)
    o = np.argsort(x, kind="mergesort")
    cum = np.cumsum(w[o]) - 0.5 * w[o]
    return float(np.interp(0.5 * w.sum(), cum, x[o]))


def conventions(data):
    out = []
    for year, d in data.items():
        for group in GROUPS:
            g = d[np.asarray(headline(group)(d))]
            pooled_mean = weighted_mean(g["NETWORTH"], g["WGT"])
            pooled_median = weighted_median(g["NETWORTH"], g["WGT"])
            out.append(dict(
                survey_year=year, education_group=group,
                mean_pooled=pooled_mean, mean_implicate_average=implicate_average(g, weighted_mean),
                median_pooled=pooled_median,
                median_implicate_average=implicate_average(g, weighted_median),
                median_pooled_interpolated=interpolated_median(g["NETWORTH"], g["WGT"]),
            ))
    o = pd.DataFrame(out)
    o["mean_diff_pct"] = 100 * (o["mean_implicate_average"] / o["mean_pooled"] - 1)
    o["median_diff_pct"] = 100 * (o["median_implicate_average"] / o["median_pooled"] - 1)
    return o.round(4)


def replicate_weights(year):
    with zipfile.ZipFile(RAW / f"scf{year}rw1s.zip") as z:
        with z.open(z.namelist()[0]) as f:
            rw = pd.read_stata(f)
    rw.columns = [c.lower() for c in rw.columns]
    # Columns are not stored in numeric order, so select them by name.
    wt = rw[[f"wt1b{i}" for i in range(1, 1000)]].to_numpy(float)
    mm = rw[[f"mm{i}" for i in range(1, 1000)]].to_numpy(float)
    w = np.nan_to_num(np.maximum(wt, 0)) * np.nan_to_num(np.maximum(mm, 0))
    return pd.DataFrame(w, index=rw["y1"].astype(int).to_numpy())


def replicate_stats(x, W):
    """Weighted mean and Fed-rule median of x for each column of weights W."""
    order = np.argsort(x, kind="mergesort")
    x, W = x[order], W[order]
    means = (W * x[:, None]).sum(0) / W.sum(0)
    cum = np.cumsum(W, axis=0)
    medians = x[np.argmax(cum >= 0.5 * cum[-1], axis=0)]
    return means, medians


def uncertainty(d, rw):
    """MEANIT-style standard errors for the latest wave."""
    imp1 = d[d["implicate"] == 1]
    if not set(imp1["record_id"]) <= set(rw.index):
        raise ValueError("replicate weights do not cover every implicate-1 record")
    W_all = rw.loc[imp1["record_id"]].to_numpy()
    totals = W_all.sum(0)
    targets = {"All families": TABLE4_ROWS["All families"],
               "College degree (all families)": TABLE4_ROWS["College degree"],
               "BA or more, with children": headline("BA or more"),
               "HS or less, with children": headline("HS or less")}
    est, imps, reps = {}, {}, {}
    for name, f in targets.items():
        g = d[np.asarray(f(d))]
        m1 = np.asarray(f(imp1))
        est[name] = (weighted_mean(g["NETWORTH"], g["WGT"]), weighted_median(g["NETWORTH"], g["WGT"]))
        imps[name] = np.array([(weighted_mean(t["NETWORTH"], t["WGT"]), weighted_median(t["NETWORTH"], t["WGT"]))
                               for _, t in g.groupby("implicate")])
        reps[name] = np.vstack(replicate_stats(imp1.loc[m1, "NETWORTH"].to_numpy(float), W_all[m1]))
    out = []

    def add(name, stat, point, imp_vals, rep_vals):
        iv = np.var(imp_vals, ddof=1)
        sv = np.sum((rep_vals - rep_vals.mean()) ** 2) / (len(rep_vals) - 1)
        se = np.sqrt(1.2 * iv + sv)
        out.append(dict(estimate=name, statistic=stat, point=point, se=se,
                        ci95_low=point - 1.96 * se, ci95_high=point + 1.96 * se,
                        se_imputation_part=np.sqrt(1.2 * iv), se_sampling_part=np.sqrt(sv)))

    for name in targets:
        for k, stat in enumerate(["mean", "median"]):
            add(name, stat, est[name][k], imps[name][:, k], reps[name][k])
    hi, lo = "BA or more, with children", "HS or less, with children"
    for k, stat in enumerate(["mean", "median"]):
        add("Gap: BA or more minus HS or less (with children)", stat,
            est[hi][k] - est[lo][k], imps[hi][:, k] - imps[lo][:, k], reps[hi][k] - reps[lo][k])
        add("Ratio: BA or more / HS or less (with children)", stat,
            est[hi][k] / est[lo][k], imps[hi][:, k] / imps[lo][:, k], reps[hi][k] / reps[lo][k])
    o = pd.DataFrame(out)
    o.attrs["replicate_total_range_millions"] = (totals.min() / 1e6, totals.max() / 1e6)
    return o


def main():
    years = survey_years()
    data = {y: load_year(y, extra=["FAMSTRUCT"]) for y in years}
    PROCESSED.mkdir(parents=True, exist_ok=True)

    cov = coverage(data)
    cov.to_csv(PROCESSED / "validation_coverage.csv", index=False)
    print("1. Coverage\n" + cov.to_string(index=False))

    t4 = fed_table4(data)
    t4.to_csv(PROCESSED / "validation_fed_table4.csv", index=False)
    print(f"\n2. Fed public-data Table 4: {len(t4)} cells compared")
    print(t4.groupby("statistic")[["abs_diff_thousands", "rel_diff"]].max().to_string())
    picks = t4[(t4["survey_year"] == max(years)) & (
        ((t4["row"] == "All families") & (t4["statistic"] == "median")) |
        ((t4["row"] == "College degree") & (t4["statistic"] == "mean")) |
        ((t4["row"] == "Couple with child(ren)") & (t4["statistic"] == "median")))]
    print("Spot-check cells:\n" + picks[["row", "statistic", "fed_cell", "fed_thousands", "ours_thousands"]].to_string(index=False))

    conv = conventions(data)
    conv.to_csv(PROCESSED / "validation_conventions.csv", index=False)
    print("\n3. Pooled vs implicate-average (percent difference, max absolute):")
    print(conv[["mean_diff_pct", "median_diff_pct"]].abs().max().to_string())
    print("   pooled vs interpolated median, max abs % diff:",
          round(float((100 * (conv["median_pooled_interpolated"] / conv["median_pooled"] - 1)).abs().max()), 3))

    latest = max(years)
    unc = uncertainty(data[latest], replicate_weights(latest))
    unc.to_csv(PROCESSED / f"validation_uncertainty_{latest}.csv", index=False)
    lo, hi = unc.attrs["replicate_total_range_millions"]
    print(f"\n4. Uncertainty, {latest} (replicate weight totals {lo:.2f}-{hi:.2f} million families)")
    show = unc.copy()
    dollars = ~show["estimate"].str.startswith("Ratio")
    for c in ["point", "se", "ci95_low", "ci95_high", "se_imputation_part", "se_sampling_part"]:
        show[c] = np.where(dollars, (show[c] / 1e3).round(1), show[c].round(2))
    print("(dollar rows in thousands of 2025 dollars)\n" + show.to_string(index=False))
    print(f"\nWrote validation CSVs to {PROCESSED.relative_to(ROOT)}/")


if __name__ == "__main__":
    main()
